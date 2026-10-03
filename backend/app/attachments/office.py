"""PDF·Word·PowerPoint 에서 글자 뽑기."""

import io

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

PDF_MAX_PAGES = 20


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = reader.pages[:PDF_MAX_PAGES]
    return "\n".join((page.extract_text() or "") for page in pages)


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_pptx(data: bytes) -> str:
    deck = Presentation(io.BytesIO(data))
    lines = []
    for number, slide in enumerate(deck.slides, start=1):
        lines.append(f"[슬라이드 {number}]")
        lines += [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    return "\n".join(lines)
